"""Safe, layout-preserving preview extraction for PowerPoint decks."""

from __future__ import annotations

import base64
from itertools import islice
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class PptxPreviewError(ValueError):
    """Raised when a presentation cannot be converted into a safe preview."""


_MAX_SLIDES = 100
_MAX_SHAPES = 500
_MAX_IMAGE_BYTES = 1024 * 1024
_MAX_EMBEDDED_IMAGE_BYTES = 5 * 1024 * 1024
_IMAGE_MIME = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "gif": "image/gif",
    "webp": "image/webp",
}


def _geometry(shape: Any, slide_width: int, slide_height: int) -> dict[str, float]:
    return {
        "x": max(0.0, min(1.0, float(shape.left) / slide_width)),
        "y": max(0.0, min(1.0, float(shape.top) / slide_height)),
        "width": max(0.0, min(1.0, float(shape.width) / slide_width)),
        "height": max(0.0, min(1.0, float(shape.height) / slide_height)),
    }


def build_pptx_preview(path: Path) -> dict[str, Any]:
    """Extract bounded slide geometry, text and browser-safe embedded images."""
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise PptxPreviewError("Could not preview presentation") from exc

    width = int(presentation.slide_width)
    height = int(presentation.slide_height)
    if width <= 0 or height <= 0:
        raise PptxPreviewError("Could not preview presentation dimensions")

    slides: list[dict[str, Any]] = []
    embedded_bytes = 0
    shape_count = 0
    for number, slide in enumerate(islice(presentation.slides, _MAX_SLIDES), start=1):
        rendered_shapes: list[dict[str, Any]] = []
        for shape in slide.shapes:
            if shape_count >= _MAX_SHAPES:
                break
            shape_count += 1
            geometry = _geometry(shape, width, height)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blob = shape.image.blob
                ext = str(shape.image.ext or "").lower()
                item: dict[str, Any] = {"kind": "image", **geometry}
                mime = _IMAGE_MIME.get(ext)
                if (
                    mime
                    and len(blob) <= _MAX_IMAGE_BYTES
                    and embedded_bytes + len(blob) <= _MAX_EMBEDDED_IMAGE_BYTES
                ):
                    item["data_url"] = (
                        f"data:{mime};base64,{base64.b64encode(blob).decode('ascii')}"
                    )
                    embedded_bytes += len(blob)
                rendered_shapes.append(item)
                continue

            if getattr(shape, "has_text_frame", False):
                text = str(getattr(shape, "text", "") or "").strip()
                if text:
                    rendered_shapes.append({"kind": "text", "text": text, **geometry})

        slides.append({"number": number, "shapes": rendered_shapes})

    return {
        "filename": path.name,
        "width": width,
        "height": height,
        "slide_count": len(presentation.slides),
        "truncated": len(presentation.slides) > _MAX_SLIDES or shape_count >= _MAX_SHAPES,
        "slides": slides,
    }
