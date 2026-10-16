"""Portable exports for saved Co-Writer documents."""

from __future__ import annotations

from io import BytesIO
import re

from .storage import CoWriterDocument


def safe_document_filename(title: str, extension: str) -> str:
    stem = re.sub(r"[^\w\-.\u4e00-\u9fff]+", "-", title.strip(), flags=re.UNICODE)
    stem = stem.strip("-._")[:100] or "cognivoy-document"
    return f"{stem}.{extension.lstrip('.')}"


def export_markdown(document: CoWriterDocument) -> bytes:
    title = document.title.strip() or "Untitled document"
    content = document.content.strip()
    body = content if content.startswith("# ") else f"# {title}\n\n{content}"
    return (body.rstrip() + "\n").encode("utf-8")


def _markdown_lines(content: str):
    in_code = False
    for raw in content.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            yield "code", line
        elif match := re.match(r"^(#{1,6})\s+(.+)$", line):
            yield f"h{min(len(match.group(1)), 3)}", match.group(2).strip()
        elif match := re.match(r"^\s*[-*+]\s+(.+)$", line):
            yield "bullet", match.group(1).strip()
        elif match := re.match(r"^\s*\d+[.)]\s+(.+)$", line):
            yield "number", match.group(1).strip()
        elif line.strip():
            yield "body", re.sub(r"[*_`]+", "", line.strip())
        else:
            yield "blank", ""


def export_docx(document: CoWriterDocument) -> bytes:
    from docx import Document
    from docx.shared import Pt

    output = BytesIO()
    target = Document()
    normal = target.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    target.add_heading(document.title.strip() or "Untitled document", level=0)
    for kind, text in _markdown_lines(document.content):
        if kind.startswith("h"):
            target.add_heading(text, level=int(kind[1]))
        elif kind == "bullet":
            target.add_paragraph(text, style="List Bullet")
        elif kind == "number":
            target.add_paragraph(text, style="List Number")
        elif kind == "code":
            paragraph = target.add_paragraph()
            run = paragraph.add_run(text)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif kind == "blank":
            target.add_paragraph()
        else:
            target.add_paragraph(text)
    target.save(output)
    return output.getvalue()


def export_pdf(document: CoWriterDocument) -> bytes:
    # ``escape`` only encodes text for ReportLab markup; it does not parse XML.
    from xml.sax.saxutils import escape  # nosec B406

    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer

    output = BytesIO()
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "CognivoyBody",
        parent=styles["BodyText"],
        fontName="STSong-Light",
        fontSize=11,
        leading=18,
        spaceAfter=7,
        wordWrap="CJK",
    )
    headings = {
        "h1": ParagraphStyle(
            "CognivoyH1", parent=base, fontSize=20, leading=27, spaceBefore=16, spaceAfter=10
        ),
        "h2": ParagraphStyle(
            "CognivoyH2", parent=base, fontSize=16, leading=23, spaceBefore=13, spaceAfter=8
        ),
        "h3": ParagraphStyle(
            "CognivoyH3", parent=base, fontSize=13, leading=20, spaceBefore=10, spaceAfter=6
        ),
    }
    title_style = ParagraphStyle(
        "CognivoyTitle",
        parent=base,
        fontSize=25,
        leading=34,
        alignment=TA_CENTER,
        spaceAfter=16,
    )
    bullet = ParagraphStyle("CognivoyBullet", parent=base, leftIndent=14, firstLineIndent=-9)
    code = ParagraphStyle(
        "CognivoyCode",
        parent=base,
        fontName="Courier",
        fontSize=8.5,
        leading=12,
        leftIndent=8,
        rightIndent=8,
    )
    story = [
        Paragraph(escape(document.title.strip() or "Untitled document"), title_style),
        Spacer(1, 5 * mm),
    ]
    for kind, text in _markdown_lines(document.content):
        escaped = escape(text)
        if kind in headings:
            story.append(Paragraph(escaped, headings[kind]))
        elif kind == "bullet":
            story.append(Paragraph(f"• {escaped}", bullet))
        elif kind == "number":
            story.append(Paragraph(escaped, bullet))
        elif kind == "code":
            story.append(Preformatted(text, code))
        elif kind == "blank":
            story.append(Spacer(1, 3 * mm))
        else:
            story.append(Paragraph(escaped, base))
    target = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=document.title,
        author="Cognivoy",
    )
    target.build(story or [PageBreak()])
    return output.getvalue()


__all__ = ["export_docx", "export_markdown", "export_pdf", "safe_document_filename"]
