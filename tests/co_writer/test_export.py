from __future__ import annotations

from io import BytesIO

from docx import Document

from cognivoy.co_writer.export import (
    export_docx,
    export_markdown,
    export_pdf,
    safe_document_filename,
)
from cognivoy.co_writer.storage import CoWriterDocument


def _document() -> CoWriterDocument:
    return CoWriterDocument(
        id="abc123def456",
        title="AI 智能体开发 / 草稿",
        content="## 架构\n\n- 规划器\n- 工具层\n\n```python\nprint('ok')\n```",
    )


def test_markdown_export_is_utf8_and_keeps_a_single_document_title() -> None:
    content = export_markdown(_document()).decode("utf-8")
    assert content.startswith("# AI 智能体开发 / 草稿\n")
    assert "## 架构" in content
    assert "print('ok')" in content


def test_docx_export_is_openable_and_preserves_content() -> None:
    payload = export_docx(_document())
    assert payload.startswith(b"PK")
    parsed = Document(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in parsed.paragraphs)
    assert "AI 智能体开发" in text
    assert "规划器" in text
    assert "print('ok')" in text


def test_pdf_export_is_a_real_selectable_text_document() -> None:
    payload = export_pdf(_document())
    assert payload.startswith(b"%PDF-")
    assert len(payload) > 1000


def test_export_filename_removes_path_separators() -> None:
    assert safe_document_filename("../课程 / 草稿", "pdf") == "课程-草稿.pdf"
