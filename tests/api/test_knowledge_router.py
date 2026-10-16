from __future__ import annotations

import asyncio
import importlib
import json
from pathlib import Path
import threading

import pytest

try:
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
except Exception:  # pragma: no cover - optional dependency in lightweight envs
    FastAPI = None
    TestClient = None

pytestmark = pytest.mark.skipif(
    FastAPI is None or TestClient is None, reason="fastapi not installed"
)

if FastAPI is not None and TestClient is not None:
    knowledge_router_module = importlib.import_module("cognivoy.api.routers.knowledge")
    router = knowledge_router_module.router
else:  # pragma: no cover - optional dependency in lightweight envs
    knowledge_router_module = None
    router = None

_REAL_ASSERT_INDEX_PREFLIGHT = (
    knowledge_router_module._assert_index_preflight if knowledge_router_module is not None else None
)


@pytest.fixture(autouse=True)
def _assume_index_engine_ready(monkeypatch):
    """Keep route tests focused unless they explicitly exercise preflight."""
    if knowledge_router_module is not None:
        monkeypatch.setattr(
            knowledge_router_module, "_assert_index_preflight", lambda _provider: None
        )


def _build_app() -> FastAPI:
    if FastAPI is None or router is None:  # pragma: no cover - guarded by pytestmark
        raise RuntimeError("fastapi is not installed")
    app = FastAPI()
    app.include_router(router, prefix="/api/v1/knowledge")
    return app


class _FakeKBManager:
    def __init__(self, base_dir: Path) -> None:
        self.base_dir = base_dir
        self.base_dir.mkdir(parents=True, exist_ok=True)
        self.config_file = self.base_dir / "kb_config.json"
        self.config: dict[str, dict] = {"knowledge_bases": {}}

    def _load_config(self) -> dict:
        return self.config

    def _save_config(self) -> None:
        pass

    def list_knowledge_bases(self) -> list[str]:
        return sorted(self.config.get("knowledge_bases", {}).keys())

    def update_kb_status(self, name: str, status: str, progress: dict | None = None) -> None:
        entry = self.config.setdefault("knowledge_bases", {}).setdefault(name, {"path": name})
        entry["status"] = status
        entry["progress"] = progress or {}

    def get_default(self, *, available_names: list[str] | None = None) -> str | None:
        names = available_names if available_names is not None else self.list_knowledge_bases()
        return names[0] if names else None

    def get_knowledge_base_path(self, name: str) -> Path:
        kb_dir = self.base_dir / name
        kb_dir.mkdir(parents=True, exist_ok=True)
        return kb_dir

    def delete_knowledge_base(self, name: str, *, confirm: bool = False) -> bool:
        if not confirm or name not in self.config.get("knowledge_bases", {}):
            return False
        self.config["knowledge_bases"].pop(name, None)
        return True

    def register_lightrag_server_kb(
        self,
        name: str,
        server_url: str,
        *,
        api_key: str = "",
        search_mode: str = "",
        description: str = "",
    ) -> dict:
        if name in self.config.get("knowledge_bases", {}):
            raise ValueError(f"A knowledge base named '{name}' already exists.")
        entry = {
            "path": name,
            "type": "lightrag_server",
            "rag_provider": "lightrag-server",
            "server_url": server_url,
            "api_key": api_key,
            "status": "ready",
        }
        if search_mode:
            entry["search_mode"] = search_mode
        self.config.setdefault("knowledge_bases", {})[name] = entry
        return entry


class _FakeInitializer:
    def __init__(self, kb_name: str, base_dir: str, **_kwargs) -> None:
        self.kb_name = kb_name
        self.base_dir = base_dir
        self.kb_dir = Path(base_dir) / kb_name
        self.raw_dir = self.kb_dir / "raw"
        self.progress_tracker = _kwargs.get("progress_tracker")

    def create_directory_structure(self) -> None:
        self.raw_dir.mkdir(parents=True, exist_ok=True)

    def _register_to_config(self) -> None:
        pass


def _upload_payload() -> list[tuple[str, tuple[str, bytes, str]]]:
    return [("files", ("demo.txt", b"hello", "text/plain"))]


def _invalid_upload_payload() -> list[tuple[str, tuple[str, bytes, str]]]:
    return [("files", ("archive.zip", b"PK\x03\x04", "application/zip"))]


def _uppercase_upload_payload() -> list[tuple[str, tuple[str, bytes, str]]]:
    return [("files", ("报告.PDF", b"%PDF-1.4\n", "application/pdf"))]


def _write_ready_llamaindex_version(kb_dir: Path) -> None:
    version_dir = kb_dir / "version-1"
    version_dir.mkdir(parents=True, exist_ok=True)
    (version_dir / "docstore.json").write_text("{}", encoding="utf-8")
    (version_dir / "index_store.json").write_text("{}", encoding="utf-8")
    (version_dir / "meta.json").write_text(
        json.dumps({"provider": "llamaindex", "signature": "sig", "version": "version-1"}),
        encoding="utf-8",
    )


def test_rag_providers_lists_llamaindex_and_pageindex(monkeypatch) -> None:
    preflight_module = importlib.import_module("cognivoy.services.rag.preflight")
    monkeypatch.setattr(
        preflight_module,
        "engine_preflight",
        lambda provider: {
            "ok": provider != "llamaindex",
            "checks": []
            if provider != "llamaindex"
            else [
                {
                    "key": "embedding",
                    "label": "Active embedding model",
                    "ok": False,
                    "detail": "Configure one under Settings → Embedding.",
                    "optional": False,
                }
            ],
        },
    )
    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/rag-providers")

    assert response.status_code == 200
    payload = response.json()
    by_id = {p["id"]: p for p in payload["providers"]}
    assert set(by_id) == {
        "llamaindex",
        "pageindex",
        "graphrag",
        "lightrag",
        "lightrag-server",
        "ima",
    }
    # LlamaIndex works out of the box; PageIndex needs an API key; GraphRAG and
    # LightRAG are optional local engines (no API key, configured = installed).
    assert by_id["llamaindex"]["requires_api_key"] is False
    assert by_id["llamaindex"]["ready"] is False
    assert by_id["llamaindex"]["preflight"]["checks"][0]["key"] == "embedding"
    assert by_id["pageindex"]["requires_api_key"] is True
    assert by_id["graphrag"]["requires_api_key"] is False
    assert by_id["lightrag"]["requires_api_key"] is False
    # LightRAG Server is a thin HTTP client: always available, no API key gate
    # (the per-KB endpoint is configured at connect time).
    assert by_id["lightrag-server"]["requires_api_key"] is False
    assert by_id["lightrag-server"]["configured"] is True
    # Same for IMA: a thin HTTPS client, credentials bound per-KB at connect
    # time rather than gated by one global key.
    assert by_id["ima"]["requires_api_key"] is False
    assert by_id["ima"]["configured"] is True
    # Mode-aware engines advertise their retrieval modes; vector engines don't.
    assert "hybrid" in by_id["lightrag"]["modes"]
    assert "mix" in by_id["lightrag-server"]["modes"]
    assert not by_id["llamaindex"].get("modes")
    # IMA exposes a single retrieval call, so it advertises no modes.
    assert not by_id["ima"].get("modes")


def test_retrieval_health_reports_actionable_index_readiness(tmp_path, monkeypatch) -> None:
    manager = _FakeKBManager(tmp_path / "kb")
    manager.config["knowledge_bases"]["healthy"] = {
        "status": "ready",
        "rag_provider": "llamaindex",
        "statistics": {
            "raw_documents": 2,
            "rag_provider": "llamaindex",
            "rag_initialized": True,
            "active_match": True,
            "needs_reindex": False,
            "index_versions": [{"ready": True}],
        },
    }
    manager.get_info = lambda name, **_kwargs: {
        "name": name,
        "status": manager.config["knowledge_bases"][name]["status"],
        "progress": None,
        "metadata": {},
        "statistics": manager.config["knowledge_bases"][name]["statistics"],
    }
    monkeypatch.setattr(knowledge_router_module, "kb_manager", manager)
    preflight_module = importlib.import_module("cognivoy.services.rag.preflight")
    monkeypatch.setattr(
        preflight_module,
        "engine_preflight",
        lambda _provider: {"ok": True, "checks": []},
    )

    with TestClient(_build_app()) as client:
        healthy = client.get("/api/v1/knowledge/healthy/retrieval-health")
        assert healthy.status_code == 200
        payload = healthy.json()
        assert payload["ready"] is True
        assert {check["key"] for check in payload["checks"]} == {
            "source_documents",
            "engine_ready",
            "active_index",
            "index_freshness",
            "task_state",
        }
        assert payload["checks"][0]["parameters"] == {"count": 2}

        manager.config["knowledge_bases"]["healthy"]["statistics"]["needs_reindex"] = True
        stale = client.get("/api/v1/knowledge/healthy/retrieval-health").json()
        assert stale["ready"] is False
        freshness = next(check for check in stale["checks"] if check["key"] == "index_freshness")
        assert freshness["action"]["label"] == "Re-index"


def test_retrieval_health_records_stable_learning_project_outcomes(tmp_path, monkeypatch) -> None:
    manager = _FakeKBManager(tmp_path / "kb")
    manager.config["knowledge_bases"]["course-notes"] = {
        "status": "ready",
        "rag_provider": "llamaindex",
        "statistics": {
            "raw_documents": 1,
            "rag_initialized": True,
            "active_match": True,
            "needs_reindex": False,
        },
    }
    manager.get_info = lambda name, **_kwargs: {
        "id": "user:kb:course-notes",
        "name": name,
        "status": manager.config["knowledge_bases"][name]["status"],
        "statistics": manager.config["knowledge_bases"][name]["statistics"],
    }
    monkeypatch.setattr(knowledge_router_module, "kb_manager", manager)
    preflight_module = importlib.import_module("cognivoy.services.rag.preflight")
    monkeypatch.setattr(
        preflight_module,
        "engine_preflight",
        lambda _provider: {"ok": True, "checks": []},
    )
    recorded: list[dict] = []

    class _ProjectStore:
        def record_linked_resource_activity(self, **kwargs):
            recorded.append(kwargs)
            return []

    project_storage = importlib.import_module("cognivoy.learning.project_storage")
    monkeypatch.setattr(project_storage, "LearningProjectStore", _ProjectStore)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/course-notes/retrieval-health")

    assert response.status_code == 200
    assert len(recorded) == 1
    assert recorded[0]["resource_type"] == "knowledge_base"
    assert "course-notes" in recorded[0]["resource_ids"]
    assert "user:kb:course-notes" in recorded[0]["resource_ids"]
    assert recorded[0]["dedupe_key"].endswith(":ready")


def test_retrieval_quality_probe_samples_source_and_scores_recall(tmp_path, monkeypatch) -> None:
    manager = _FakeKBManager(tmp_path / "kb")
    manager.config["knowledge_bases"]["course-notes"] = {
        "status": "ready",
        "rag_provider": "llamaindex",
    }
    raw_dir = manager.get_knowledge_base_path("course-notes") / "raw"
    raw_dir.mkdir(parents=True)
    source_text = (
        "Retrieval augmented generation combines external evidence with model reasoning "
        "so answers can cite a controlled source instead of relying only on parameters."
    )
    (raw_dir / "lesson.txt").write_text(source_text, encoding="utf-8")
    monkeypatch.setattr(knowledge_router_module, "kb_manager", manager)

    rag_service = importlib.import_module("cognivoy.services.rag.service")

    class _FakeRAGService:
        def __init__(self, **_kwargs) -> None:
            pass

        async def search(self, query, kb_name, **_kwargs):
            assert kb_name == "course-notes"
            assert "Retrieval augmented generation" in query
            return {
                "content": source_text,
                "sources": [
                    {
                        "source": str(raw_dir / "lesson.txt"),
                        "title": "lesson.txt",
                        "content": source_text,
                    }
                ],
            }

    monkeypatch.setattr(rag_service, "RAGService", _FakeRAGService)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/course-notes/retrieval-quality")

    assert response.status_code == 200
    payload = response.json()
    assert payload["quality"] == "good"
    assert payload["score"] == 1.0
    assert payload["sample_source"] == "lesson.txt"
    assert payload["source_matched"] is True
    assert payload["retrieved_count"] == 1


def test_index_fragment_workbench_returns_bounded_real_hits(tmp_path, monkeypatch) -> None:
    manager = _FakeKBManager(tmp_path / "kb")
    manager.config["knowledge_bases"]["course-notes"] = {
        "status": "ready",
        "rag_provider": "llamaindex",
    }
    monkeypatch.setattr(knowledge_router_module, "kb_manager", manager)
    rag_service = importlib.import_module("cognivoy.services.rag.service")

    class _FakeRAGService:
        def __init__(self, **_kwargs) -> None:
            pass

        async def search(self, query, kb_name, **kwargs):
            assert query == "agent memory"
            assert kb_name == "course-notes"
            assert kwargs["top_k"] == 3
            return {
                "sources": [
                    {
                        "source": "chapter-4.md",
                        "page": 12,
                        "score": 0.91,
                        "content": "Long-term memory needs explicit provenance.",
                        "embedding": [0.1, 0.2],
                    }
                ]
            }

    monkeypatch.setattr(rag_service, "RAGService", _FakeRAGService)
    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/course-notes/index-fragments",
            json={"query": "agent memory", "top_k": 3},
        )

    assert response.status_code == 200
    fragment = response.json()["fragments"][0]
    assert fragment["source"] == "chapter-4.md"
    assert fragment["content"].startswith("Long-term memory")
    assert "embedding" not in fragment["metadata"]


def test_set_rag_provider_mode_persists_validates_and_reflects() -> None:
    with TestClient(_build_app()) as client:
        ok = client.put("/api/v1/knowledge/rag-providers/lightrag/mode", json={"mode": "MIX"})
        assert ok.status_code == 200
        assert ok.json()["mode"] == "mix"  # normalized

        providers = client.get("/api/v1/knowledge/rag-providers").json()["providers"]
        by_id = {p["id"]: p for p in providers}
        assert by_id["lightrag"]["default_mode"] == "mix"

        # Invalid mode for the engine → 400; mode-less engine → 404.
        assert (
            client.put(
                "/api/v1/knowledge/rag-providers/lightrag/mode", json={"mode": "bogus"}
            ).status_code
            == 400
        )
        assert (
            client.put(
                "/api/v1/knowledge/rag-providers/llamaindex/mode", json={"mode": "x"}
            ).status_code
            == 404
        )


def test_supported_file_types_returns_upload_policy() -> None:
    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/supported-file-types")

    assert response.status_code == 200
    payload = response.json()
    assert ".pdf" in payload["extensions"]
    assert ".docx" in payload["extensions"]
    assert ".xlsx" in payload["extensions"]
    assert ".pptx" in payload["extensions"]
    assert ".md" in payload["extensions"]
    assert ".png" in payload["extensions"]
    assert payload["max_file_size_bytes"] > 0
    assert "max_pdf_size_bytes" not in payload
    assert ".pdf" in payload["accept"]
    assert ".docx" in payload["accept"]
    assert ".png" in payload["accept"]
    assert "image/png" in payload["accept"]


def test_create_kb_does_not_require_llm_precheck(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "KnowledgeBaseInitializer", _FakeInitializer)
    monkeypatch.setattr(
        knowledge_router_module,
        "get_llm_config",
        lambda: (_ for _ in ()).throw(RuntimeError("should not be called")),
        raising=False,
    )

    async def _noop_init_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_initialization_task", _noop_init_task)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/create",
            data={"name": "kb-new", "rag_provider": "llamaindex"},
            files=_upload_payload(),
        )

    assert response.status_code == 200
    body = response.json()
    assert body["name"] == "kb-new"
    assert isinstance(body.get("task_id"), str) and body["task_id"]
    assert manager.config["knowledge_bases"]["kb-new"]["rag_provider"] == "llamaindex"
    assert manager.config["knowledge_bases"]["kb-new"]["needs_reindex"] is False


def test_create_rejects_missing_embedding_before_registering_kb(
    monkeypatch, tmp_path: Path
) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(
        knowledge_router_module,
        "_assert_index_preflight",
        _REAL_ASSERT_INDEX_PREFLIGHT,
    )

    preflight_module = importlib.import_module("cognivoy.services.rag.preflight")
    monkeypatch.setattr(
        preflight_module,
        "engine_preflight",
        lambda _provider: {
            "ok": False,
            "checks": [
                {
                    "key": "embedding",
                    "label": "Active embedding model",
                    "ok": False,
                    "detail": "Configure one under Settings → Embedding.",
                    "optional": False,
                }
            ],
        },
    )

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/create",
            data={"name": "kb-no-embedding", "rag_provider": "llamaindex"},
            files=_upload_payload(),
        )

    assert response.status_code == 422
    detail = response.json()["detail"]
    assert detail["code"] == "embedding_model_missing"
    assert detail["retryable"] is False
    assert detail["action"]["path"] == "/settings/embedding"
    assert detail["checks"][0]["key"] == "embedding"
    assert manager.config["knowledge_bases"] == {}


def test_create_coerces_legacy_provider_to_llamaindex(monkeypatch, tmp_path: Path) -> None:
    """Unknown/removed provider strings silently normalize to llamaindex."""
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    async def _noop_init_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_initialization_task", _noop_init_task)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/create",
            data={"name": "kb-legacy", "rag_provider": "raganything"},
            files=_upload_payload(),
        )

    assert response.status_code == 200
    assert manager.config["knowledge_bases"]["kb-legacy"]["rag_provider"] == "llamaindex"


def test_create_preserves_known_nondefault_provider(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "KnowledgeBaseInitializer", _FakeInitializer)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    pageindex_config = importlib.import_module("cognivoy.services.rag.pipelines.pageindex.config")
    monkeypatch.setattr(pageindex_config, "is_pageindex_configured", lambda: True)

    async def _noop_init_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_initialization_task", _noop_init_task)

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/create",
            data={"name": "kb-page", "rag_provider": "pageindex"},
            files=[("files", ("demo.pdf", b"%PDF-1.4\n", "application/pdf"))],
        )

    assert response.status_code == 200
    assert manager.config["knowledge_bases"]["kb-page"]["rag_provider"] == "pageindex"


def test_create_rejects_invalid_files_before_registering_kb(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/create",
            data={"name": "kb-invalid", "rag_provider": "llamaindex"},
            files=_invalid_upload_payload(),
        )

    assert response.status_code == 400
    assert "unsupported file type" in response.json()["detail"].lower()
    assert "kb-invalid" not in manager.config["knowledge_bases"]


def test_create_rejects_invalid_kb_name_before_registering_kb(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/create",
            data={"name": "bad/name", "rag_provider": "llamaindex"},
            files=_upload_payload(),
        )

    assert response.status_code == 400
    assert "reserved characters" in response.json()["detail"].lower()
    assert manager.config["knowledge_bases"] == {}


def test_create_normalizes_uploaded_extension_to_lowercase(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "KnowledgeBaseInitializer", _FakeInitializer)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    async def _noop_init_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_initialization_task", _noop_init_task)

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/create",
            data={"name": "kb-uppercase", "rag_provider": "llamaindex"},
            files=_uppercase_upload_payload(),
        )

    assert response.status_code == 200
    assert response.json()["files"] == ["报告.pdf"]
    assert (tmp_path / "knowledge_bases" / "kb-uppercase" / "raw" / "报告.pdf").exists()


def test_upload_returns_409_when_kb_needs_reindex(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["legacy-kb"] = {
        "path": "legacy-kb",
        "rag_provider": "llamaindex",
        "needs_reindex": True,
        "status": "needs_reindex",
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/legacy-kb/upload", files=_upload_payload())

    assert response.status_code == 409
    assert "needs reindex" in response.json()["detail"].lower()


def test_upload_ready_kb_returns_task_id(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["ready-kb"] = {
        "path": "ready-kb",
        "rag_provider": "llamaindex",
        "needs_reindex": False,
        "status": "ready",
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    async def _noop_upload_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_upload_processing_task", _noop_upload_task)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/ready-kb/upload", files=_upload_payload())

    assert response.status_code == 200
    body = response.json()
    assert isinstance(body.get("task_id"), str) and body["task_id"]


def test_upload_preflight_detects_same_path_and_identical_content(
    monkeypatch, tmp_path: Path
) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["ready-kb"] = {
        "path": "ready-kb",
        "rag_provider": "llamaindex",
        "status": "ready",
    }
    raw_dir = manager.get_knowledge_base_path("ready-kb") / "raw"
    raw_dir.mkdir(parents=True)
    content = b"same bytes"
    (raw_dir / "existing.txt").write_bytes(content)
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/ready-kb/upload-preflight",
            json={
                "files": [
                    {
                        "name": "existing.txt",
                        "size_bytes": len(content),
                        "sha256": __import__("hashlib").sha256(content).hexdigest(),
                    },
                    {
                        "name": "renamed.txt",
                        "size_bytes": len(content),
                        "sha256": __import__("hashlib").sha256(content).hexdigest(),
                    },
                ]
            },
        )

    assert response.status_code == 200
    body = response.json()
    assert body["has_duplicates"] is True
    assert body["files"][0]["path_exists"] is True
    assert body["files"][1]["same_content_paths"] == ["existing.txt"]


def test_upload_requires_explicit_duplicate_strategy_and_can_rename(
    monkeypatch, tmp_path: Path
) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["ready-kb"] = {
        "path": "ready-kb",
        "rag_provider": "llamaindex",
        "status": "ready",
    }
    raw_dir = manager.get_knowledge_base_path("ready-kb") / "raw"
    raw_dir.mkdir(parents=True)
    (raw_dir / "demo.txt").write_text("old", encoding="utf-8")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    async def _noop_upload_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_upload_processing_task", _noop_upload_task)
    with TestClient(_build_app()) as client:
        conflict = client.post("/api/v1/knowledge/ready-kb/upload", files=_upload_payload())
        renamed = client.post(
            "/api/v1/knowledge/ready-kb/upload",
            data={"duplicate_strategy": "rename"},
            files=_upload_payload(),
        )

    assert conflict.status_code == 409
    assert conflict.json()["detail"]["code"] == "knowledge_duplicate_files"
    assert renamed.status_code == 200
    assert renamed.json()["renamed_files"] == {"demo.txt": "demo (2).txt"}
    assert (raw_dir / "demo.txt").read_text(encoding="utf-8") == "old"
    assert (raw_dir / "demo (2).txt").read_text(encoding="utf-8") == "hello"


def test_resumable_upload_restores_offset_and_starts_single_processing_task(
    monkeypatch, tmp_path: Path
) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["ready-kb"] = {
        "path": "ready-kb",
        "rag_provider": "llamaindex",
        "status": "ready",
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    async def _noop_upload_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_upload_processing_task", _noop_upload_task)
    content = b"resumable upload"
    with TestClient(_build_app()) as client:
        created = client.post(
            "/api/v1/knowledge/ready-kb/uploads/resumable",
            json={"name": "large.txt", "size_bytes": len(content)},
        )
        assert created.status_code == 200
        session_id = created.json()["session_id"]
        first = client.put(
            f"/api/v1/knowledge/ready-kb/uploads/resumable/{session_id}",
            data={"offset": "0"},
            files={"chunk": ("large.txt", content[:5], "application/octet-stream")},
        )
        restored = client.get(f"/api/v1/knowledge/ready-kb/uploads/resumable/{session_id}")
        second = client.put(
            f"/api/v1/knowledge/ready-kb/uploads/resumable/{session_id}",
            data={"offset": "5"},
            files={"chunk": ("large.txt", content[5:], "application/octet-stream")},
        )
        completed = client.post(
            f"/api/v1/knowledge/ready-kb/uploads/resumable/{session_id}/complete",
            json={"duplicate_strategy": "error"},
        )
        processed = client.post(
            "/api/v1/knowledge/ready-kb/uploads/resumable/process",
            json={"paths": ["large.txt"]},
        )

    assert first.json()["offset"] == 5
    assert restored.json()["offset"] == 5
    assert second.json()["complete"] is True
    assert completed.json() == {"session_id": session_id, "path": "large.txt", "skipped": False}
    assert isinstance(processed.json()["task_id"], str)
    assert (
        manager.get_knowledge_base_path("ready-kb") / "raw" / "large.txt"
    ).read_bytes() == content


def test_resumable_upload_cleanup_removes_only_expired_sessions(
    monkeypatch, tmp_path: Path
) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["ready-kb"] = {
        "path": "ready-kb",
        "rag_provider": "llamaindex",
        "status": "ready",
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    with TestClient(_build_app()) as client:
        expired = client.post(
            "/api/v1/knowledge/ready-kb/uploads/resumable",
            json={"name": "expired.txt", "size_bytes": 10},
        ).json()
        active = client.post(
            "/api/v1/knowledge/ready-kb/uploads/resumable",
            json={"name": "active.txt", "size_bytes": 10},
        ).json()
        sessions = manager.get_knowledge_base_path("ready-kb") / ".upload-sessions"
        expired_path = sessions / f"{expired['session_id']}.json"
        expired_payload = json.loads(expired_path.read_text(encoding="utf-8"))
        expired_payload["updated_at"] = "2020-01-01T00:00:00+00:00"
        expired_path.write_text(json.dumps(expired_payload), encoding="utf-8")

        cleaned = client.post("/api/v1/knowledge/ready-kb/uploads/resumable/cleanup")

    assert cleaned.status_code == 200
    assert cleaned.json() == {
        "removed_session_ids": [expired["session_id"]],
        "removed_count": 1,
    }
    assert not expired_path.exists()
    assert not (sessions / f"{expired['session_id']}.part").exists()
    assert (sessions / f"{active['session_id']}.json").exists()
    assert (sessions / f"{active['session_id']}.part").exists()


def test_upload_flips_ready_kb_to_processing_before_dispatch(monkeypatch, tmp_path: Path) -> None:
    """An existing ready KB must not keep reporting ``ready`` between the
    accepted upload response and the background task's first progress write."""
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["ready-kb"] = {
        "path": "ready-kb",
        "rag_provider": "llamaindex",
        "needs_reindex": False,
        "status": "ready",
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    async def _noop_upload_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_upload_processing_task", _noop_upload_task)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/ready-kb/upload", files=_upload_payload())

    assert response.status_code == 200
    entry = manager.config["knowledge_bases"]["ready-kb"]
    assert entry["status"] == "processing"
    # Stage must be a member of the frontend's LIVE_PROGRESS_STAGES set.
    assert entry["progress"]["stage"] == "starting"
    assert entry["progress"]["task_id"] == response.json()["task_id"]


def test_running_upload_can_be_cancelled_through_api(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["ready-kb"] = {
        "path": "ready-kb",
        "rag_provider": "llamaindex",
        "needs_reindex": False,
        "status": "ready",
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", manager.base_dir)

    started = threading.Event()
    stopped = threading.Event()

    async def _blocking_upload_task(*_args, **_kwargs):
        started.set()
        try:
            await asyncio.Event().wait()
        finally:
            stopped.set()

    monkeypatch.setattr(
        knowledge_router_module,
        "run_upload_processing_task",
        _blocking_upload_task,
    )

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/ready-kb/upload", files=_upload_payload())
        assert response.status_code == 200
        task_id = response.json()["task_id"]
        assert started.wait(timeout=1)

        cancel = client.post(f"/api/v1/knowledge/ready-kb/tasks/{task_id}/cancel")

        assert cancel.status_code == 200
        assert cancel.json()["status"] == "cancelled"
        assert stopped.wait(timeout=1)
        progress = knowledge_router_module.ProgressTracker(
            "ready-kb", manager.base_dir
        ).get_progress()
        assert progress is not None
        assert progress["stage"] == "cancelled"


@pytest.mark.asyncio
async def test_upload_runner_persists_cancelled_terminal_state(monkeypatch, tmp_path: Path) -> None:
    base_dir = tmp_path / "knowledge_bases"
    raw_dir = base_dir / "ready-kb" / "raw"
    raw_dir.mkdir(parents=True)
    source = raw_dir / "note.txt"
    source.write_text("hello", encoding="utf-8")
    started = asyncio.Event()

    class _BlockingAdder:
        def __init__(self, *_args, **_kwargs) -> None:
            pass

        def add_documents(self, *_args, **_kwargs):
            return [source]

        async def process_new_documents(self, *_args, **_kwargs):
            started.set()
            await asyncio.Event().wait()

    monkeypatch.setattr(knowledge_router_module, "DocumentAdder", _BlockingAdder)
    task_id = knowledge_router_module._build_unique_task_id("kb_upload", "runner-cancel-test")
    task = asyncio.create_task(
        knowledge_router_module.run_upload_processing_task(
            kb_name="ready-kb",
            base_dir=str(base_dir),
            uploaded_file_paths=[str(source)],
            task_id=task_id,
            rag_provider="llamaindex",
        )
    )
    await asyncio.wait_for(started.wait(), timeout=1)

    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task

    progress = knowledge_router_module.ProgressTracker("ready-kb", base_dir).get_progress()
    assert progress["stage"] == "cancelled"
    metadata = knowledge_router_module.TaskIDManager.get_instance().get_task_metadata(task_id)
    assert metadata["status"] == "cancelled"


def test_cancel_rejects_stale_task_id(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["ready-kb"] = {
        "path": "ready-kb",
        "status": "processing",
        "progress": {"stage": "processing_documents", "task_id": "new-task"},
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/ready-kb/tasks/old-task/cancel")

    assert response.status_code == 409
    assert "no longer the active job" in response.json()["detail"]


def test_upload_task_marks_provider_failures_as_error(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("COGNIVOY_RAG_RETRIEVAL_PROFILE", "vector")
    base_dir = tmp_path / "knowledge_bases"
    kb_dir = base_dir / "kb"
    raw_dir = kb_dir / "raw"
    raw_dir.mkdir(parents=True)
    _write_ready_llamaindex_version(kb_dir)
    (base_dir / "kb_config.json").write_text(
        json.dumps(
            {
                "knowledge_bases": {
                    "kb": {
                        "path": "kb",
                        "rag_provider": "llamaindex",
                        "status": "ready",
                    }
                }
            }
        ),
        encoding="utf-8",
    )
    source = tmp_path / "bad.txt"
    source.write_text("bad", encoding="utf-8")

    class _FailingRagService:
        def __init__(self, *_args, **_kwargs) -> None:
            pass

        async def add_documents(self, *_args, **_kwargs) -> bool:
            raise RuntimeError("parse failed loudly")

    monkeypatch.setattr(
        "cognivoy.knowledge.add_documents.RAGService",
        _FailingRagService,
    )

    asyncio.run(
        knowledge_router_module.run_upload_processing_task(
            kb_name="kb",
            base_dir=str(base_dir),
            uploaded_file_paths=[str(source)],
            task_id="upload-failure-test",
            rag_provider="llamaindex",
        )
    )

    persisted = json.loads((base_dir / "kb_config.json").read_text(encoding="utf-8"))
    entry = persisted["knowledge_bases"]["kb"]
    assert entry["status"] == "error"
    assert "parse failed loudly" not in entry["last_error"]
    assert "failed (bad.txt)" in entry["last_error"]
    assert "Diagnostic ID:" in entry["last_error"]
    assert entry["progress"]["stage"] == "error"
    assert entry["progress"]["indexed_count"] == 0

    events = list(knowledge_router_module.get_task_stream_manager()._buffers["upload-failure-test"])
    assert not any("parse failed loudly" in str(event) for event in events)
    failed = next(event["payload"] for event in events if event["event"] == "failed")
    assert failed["code"] == "knowledge_document_partial_failure"
    assert failed["retryable"] is True
    assert failed["trace_id"]
    assert failed["details"] == {}


def test_list_files_accepts_default_alias(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["actual-kb"] = {
        "path": "actual-kb",
        "status": "ready",
    }
    raw_dir = manager.base_dir / "actual-kb" / "raw"
    raw_dir.mkdir(parents=True)
    (raw_dir / "demo.txt").write_text("hello", encoding="utf-8")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/default/files")

    assert response.status_code == 200
    assert response.json()["files"][0]["name"] == "demo.txt"


def test_list_fallback_reports_error_status(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["broken-kb"] = {
        "path": "broken-kb",
        "status": "ready",
    }
    (manager.base_dir / "broken-kb").mkdir(parents=True)
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/list")

    assert response.status_code == 200
    [item] = response.json()
    assert item["status"] == "error"
    assert item["progress"]["stage"] == "error"
    assert "get_info" in item["progress"]["error"]


def test_list_reuses_manager_config_snapshot(monkeypatch, tmp_path: Path) -> None:
    class _CountingKBManager:
        def __init__(self) -> None:
            self.base_dir = tmp_path / "knowledge_bases"
            self.base_dir.mkdir(parents=True)
            self.names = ["kb-a", "kb-b", "kb-c"]
            self.list_calls = 0
            self.default_calls = 0
            self.info_calls: list[tuple[str, bool, str | None]] = []

        def list_knowledge_bases(self) -> list[str]:
            self.list_calls += 1
            return self.names

        def get_default(self, *, available_names: list[str] | None = None) -> str:
            self.default_calls += 1
            assert available_names == self.names
            return self.names[0]

        def get_info(
            self,
            name: str,
            *,
            refresh_config: bool,
            default_name: str | None,
        ) -> dict:
            self.info_calls.append((name, refresh_config, default_name))
            return {
                "name": name,
                "path": str(self.base_dir / name),
                "is_default": name == default_name,
                "statistics": {},
                "metadata": {"name": name},
                "status": "ready",
                "progress": None,
            }

    manager = _CountingKBManager()
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "list_visible_kb_access", lambda: [])

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/list")

    assert response.status_code == 200
    assert [item["name"] for item in response.json()] == manager.names
    assert manager.list_calls == 1
    assert manager.default_calls == 1
    assert manager.info_calls == [(name, False, "kb-a") for name in manager.names]


def _ready_kb_manager(tmp_path: Path, name: str = "kb") -> "_FakeKBManager":
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"][name] = {
        "path": name,
        "rag_provider": "llamaindex",
        "needs_reindex": False,
        "status": "ready",
    }
    (manager.base_dir / name / "raw").mkdir(parents=True, exist_ok=True)
    return manager


def test_delete_requires_current_impact_confirmation_and_records_tombstone(
    monkeypatch, tmp_path: Path
) -> None:
    from cognivoy.knowledge.references import KnowledgeReferenceService

    manager = _ready_kb_manager(tmp_path, "course-kb")
    references = KnowledgeReferenceService(
        db_path=tmp_path / "chat.db",
        notebook_dir=tmp_path / "workspace" / "notebook",
        workspace_dir=tmp_path / "workspace",
        tombstone_path=tmp_path / "tombstones.json",
    )
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(
        knowledge_router_module, "get_knowledge_reference_service", lambda: references
    )

    with TestClient(_build_app()) as client:
        preview = client.get("/api/v1/knowledge/course-kb/deletion-impact")
        assert preview.status_code == 200
        token = preview.json()["confirmation_token"]
        assert client.delete("/api/v1/knowledge/course-kb").status_code == 409
        deleted = client.delete(f"/api/v1/knowledge/course-kb?confirmation_token={token}")

    assert deleted.status_code == 200
    assert references.get_tombstone("course-kb")["status"] == "deleted"
    assert "course-kb" not in manager.config["knowledge_bases"]


def test_create_folder_makes_subdir(monkeypatch, tmp_path: Path) -> None:
    manager = _ready_kb_manager(tmp_path)
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/kb/folders", json={"path": "Papers/2024"})

    assert response.status_code == 200
    assert response.json()["path"] == "Papers/2024"
    assert (manager.base_dir / "kb" / "raw" / "Papers" / "2024").is_dir()


def test_create_folder_rejects_traversal(monkeypatch, tmp_path: Path) -> None:
    manager = _ready_kb_manager(tmp_path)
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/kb/folders", json={"path": "../escape"})

    assert response.status_code == 400


def test_list_files_returns_nested_tree(monkeypatch, tmp_path: Path) -> None:
    manager = _ready_kb_manager(tmp_path)
    raw = manager.base_dir / "kb" / "raw"
    (raw / "Papers").mkdir(parents=True)
    (raw / "Papers" / "a.pdf").write_text("%PDF-1.4\n", encoding="utf-8")
    (raw / "root.txt").write_text("hi", encoding="utf-8")
    (raw / "Empty").mkdir()
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/kb/files")

    assert response.status_code == 200
    entries = {e["name"]: e for e in response.json()["files"]}
    assert entries["Papers"]["type"] == "folder"
    assert entries["Empty"]["type"] == "folder"  # empty folder still shows
    assert entries["Papers/a.pdf"]["type"] == "file"
    assert entries["root.txt"]["type"] == "file"


def test_remote_kb_file_listing_is_empty_without_creating_local_storage(
    monkeypatch, tmp_path: Path
) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.register_lightrag_server_kb("remote", "http://localhost:9621")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/remote/files")

    assert response.status_code == 200
    assert response.json() == {"files": []}
    assert not (manager.base_dir / "remote").exists()


@pytest.mark.parametrize(
    ("method", "url", "kwargs"),
    [
        ("get", "/api/v1/knowledge/remote/files/demo.txt", {}),
        ("get", "/api/v1/knowledge/remote/file-preview-text/demo.txt", {}),
        ("delete", "/api/v1/knowledge/remote/files/demo.txt", {}),
        ("post", "/api/v1/knowledge/remote/folders", {"json": {"path": "notes"}}),
        (
            "post",
            "/api/v1/knowledge/remote/files/move",
            {"json": {"source": "demo.txt", "dest_folder": "notes"}},
        ),
        ("post", "/api/v1/knowledge/remote/upload", {"files": _upload_payload()}),
    ],
)
def test_remote_kb_rejects_local_file_operations_without_creating_storage(
    monkeypatch, tmp_path: Path, method: str, url: str, kwargs: dict
) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.register_lightrag_server_kb("remote", "http://localhost:9621")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = getattr(client, method)(url, **kwargs)

    assert response.status_code == 409
    assert "external resource" in response.json()["detail"]
    assert not (manager.base_dir / "remote").exists()


def test_list_files_returns_404_for_unknown_kb_without_creating_storage(
    monkeypatch, tmp_path: Path
) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/missing/files")

    assert response.status_code == 404
    assert not (manager.base_dir / "missing").exists()


def test_raw_file_download_rejects_traversal(monkeypatch, tmp_path: Path) -> None:
    manager = _ready_kb_manager(tmp_path)
    (manager.base_dir / "secret.txt").write_text("secret", encoding="utf-8")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/kb/files/%2E%2E/secret.txt")

    assert response.status_code == 403


def test_upload_preserves_folder_structure(monkeypatch, tmp_path: Path) -> None:
    manager = _ready_kb_manager(tmp_path)
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    async def _noop_upload_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_upload_processing_task", _noop_upload_task)

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/kb/upload",
            files=[("files", ("note.txt", b"hi", "text/plain"))],
            data={"rel_paths": "MyFolder/sub/note.txt"},
        )

    assert response.status_code == 200
    assert (manager.base_dir / "kb" / "raw" / "MyFolder" / "sub" / "note.txt").is_file()


def test_upload_allows_same_filename_in_different_folders(monkeypatch, tmp_path: Path) -> None:
    manager = _ready_kb_manager(tmp_path)
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    async def _noop_upload_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_upload_processing_task", _noop_upload_task)

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/kb/upload",
            files=[
                ("files", ("note.txt", b"one", "text/plain")),
                ("files", ("note.txt", b"two", "text/plain")),
            ],
            data={"rel_paths": ["ModuleA/note.txt", "ModuleB/note.txt"]},
        )

    assert response.status_code == 200
    assert (manager.base_dir / "kb" / "raw" / "ModuleA" / "note.txt").is_file()
    assert (manager.base_dir / "kb" / "raw" / "ModuleB" / "note.txt").is_file()


def test_move_file_into_folder(monkeypatch, tmp_path: Path) -> None:
    manager = _ready_kb_manager(tmp_path)
    raw = manager.base_dir / "kb" / "raw"
    (raw / "demo.txt").write_text("hi", encoding="utf-8")
    (raw / "Papers").mkdir()
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", tmp_path / "knowledge_bases")

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/kb/files/move",
            json={"source": "demo.txt", "dest_folder": "Papers"},
        )

    assert response.status_code == 200
    assert (raw / "Papers" / "demo.txt").is_file()
    assert not (raw / "demo.txt").exists()


def test_list_files_preserves_kb_named_default(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["actual-kb"] = {
        "path": "actual-kb",
        "status": "ready",
    }
    manager.config["knowledge_bases"]["default"] = {
        "path": "default",
        "status": "ready",
    }
    actual_raw = manager.base_dir / "actual-kb" / "raw"
    actual_raw.mkdir(parents=True)
    (actual_raw / "actual.txt").write_text("hello", encoding="utf-8")
    default_raw = manager.base_dir / "default" / "raw"
    default_raw.mkdir(parents=True)
    (default_raw / "default.txt").write_text("hello", encoding="utf-8")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/default/files")

    assert response.status_code == 200
    assert response.json()["files"][0]["name"] == "default.txt"


def test_file_preview_text_accepts_default_alias(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["actual-kb"] = {
        "path": "actual-kb",
        "status": "ready",
    }
    raw_dir = manager.base_dir / "actual-kb" / "raw"
    raw_dir.mkdir(parents=True)
    target = raw_dir / "slides.pptx"
    target.write_bytes(b"PK\x03\x04")
    calls: dict[str, object] = {}

    def _fake_extract(path: Path, **kwargs) -> str:
        calls["path"] = path
        calls["kwargs"] = kwargs
        return "--- Slide 1 ---\nTitle"

    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "extract_text_from_path", _fake_extract)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/default/file-preview-text/slides.pptx")

    assert response.status_code == 200
    assert response.text == "--- Slide 1 ---\nTitle"
    assert calls["path"] == target
    assert calls["kwargs"]["max_chars"] == 200_000


def test_file_preview_text_returns_422_for_extraction_errors(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["actual-kb"] = {
        "path": "actual-kb",
        "status": "ready",
    }
    raw_dir = manager.base_dir / "actual-kb" / "raw"
    raw_dir.mkdir(parents=True)
    (raw_dir / "slides.pptx").write_bytes(b"PK\x03\x04")
    extraction_error = knowledge_router_module.DocumentExtractionError

    def _fake_extract(*_args, **_kwargs) -> str:
        raise extraction_error("slides.pptx: no extractable text")

    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "extract_text_from_path", _fake_extract)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/actual-kb/file-preview-text/slides.pptx")

    assert response.status_code == 422
    assert "no extractable text" in response.json()["detail"]


def test_pptx_slide_preview_preserves_slide_geometry_text_and_images(
    monkeypatch, tmp_path: Path
) -> None:
    from base64 import b64decode

    from pptx import Presentation
    from pptx.util import Inches

    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["actual-kb"] = {
        "path": "actual-kb",
        "status": "ready",
    }
    raw_dir = manager.base_dir / "actual-kb" / "raw"
    raw_dir.mkdir(parents=True)
    image = raw_dir / "pixel.png"
    image.write_bytes(
        b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
    )
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "Architecture overview"
    slide.shapes.add_picture(str(image), Inches(6), Inches(1), width=Inches(1))
    target = raw_dir / "slides.pptx"
    presentation.save(target)
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        response = client.get("/api/v1/knowledge/default/file-preview-slides/slides.pptx")

    assert response.status_code == 200
    payload = response.json()
    assert payload["filename"] == "slides.pptx"
    assert payload["slide_count"] == 1
    assert payload["width"] > payload["height"]
    shapes = payload["slides"][0]["shapes"]
    text_shape = next(item for item in shapes if item["kind"] == "text")
    assert text_shape["text"] == "Architecture overview"
    assert 0 < text_shape["x"] < 1
    assert 0 < text_shape["width"] < 1
    image_shape = next(item for item in shapes if item["kind"] == "image")
    assert image_shape["data_url"].startswith("data:image/png;base64,")


def test_pptx_slide_preview_rejects_wrong_type_and_corrupt_deck(
    monkeypatch, tmp_path: Path
) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["actual-kb"] = {
        "path": "actual-kb",
        "status": "ready",
    }
    raw_dir = manager.base_dir / "actual-kb" / "raw"
    raw_dir.mkdir(parents=True)
    (raw_dir / "notes.txt").write_text("hello", encoding="utf-8")
    (raw_dir / "broken.pptx").write_bytes(b"not a presentation")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)

    with TestClient(_build_app()) as client:
        wrong = client.get("/api/v1/knowledge/actual-kb/file-preview-slides/notes.txt")
        corrupt = client.get("/api/v1/knowledge/actual-kb/file-preview-slides/broken.pptx")

    assert wrong.status_code == 415
    assert corrupt.status_code == 422
    assert "preview" in corrupt.json()["detail"].lower()


def test_reindex_accepts_default_alias(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("COGNIVOY_RAG_RETRIEVAL_PROFILE", "vector")
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["actual-kb"] = {
        "path": "actual-kb",
        "status": "ready",
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", manager.base_dir)

    class _Signature:
        def hash(self) -> str:
            return "sig"

    embedding_signature = importlib.import_module("cognivoy.services.rag.embedding_signature")
    index_versioning = importlib.import_module("cognivoy.services.rag.index_versioning")
    monkeypatch.setattr(
        embedding_signature, "signature_from_embedding_config", lambda: _Signature()
    )
    monkeypatch.setattr(index_versioning, "find_matching_version", lambda *_args, **_kwargs: None)

    async def _noop_reindex_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_reindex_task", _noop_reindex_task)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/default/reindex")

    assert response.status_code == 200
    body = response.json()
    assert body["noop"] is False
    assert isinstance(body.get("task_id"), str) and body["task_id"]
    assert manager.config["knowledge_bases"]["actual-kb"]["status"] == "initializing"


def test_reindex_error_status_bypasses_existing_match_noop(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("COGNIVOY_RAG_RETRIEVAL_PROFILE", "vector")
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["failed-kb"] = {
        "path": "failed-kb",
        "status": "error",
        "progress": {"stage": "error", "message": "previous indexing failed"},
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", manager.base_dir)

    class _Signature:
        def hash(self) -> str:
            return "sig"

    embedding_signature = importlib.import_module("cognivoy.services.rag.embedding_signature")
    index_versioning = importlib.import_module("cognivoy.services.rag.index_versioning")
    monkeypatch.setattr(
        embedding_signature, "signature_from_embedding_config", lambda: _Signature()
    )
    monkeypatch.setattr(
        index_versioning,
        "find_matching_version",
        lambda *_args, **_kwargs: {"layout": "flat", "ready": True},
    )

    async def _noop_reindex_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_reindex_task", _noop_reindex_task)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/failed-kb/reindex")

    assert response.status_code == 200
    body = response.json()
    assert body["noop"] is False
    assert isinstance(body.get("task_id"), str) and body["task_id"]
    assert manager.config["knowledge_bases"]["failed-kb"]["status"] == "initializing"


def test_retry_error_status_queues_reindex(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("COGNIVOY_RAG_RETRIEVAL_PROFILE", "vector")
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["failed-kb"] = {
        "path": "failed-kb",
        "status": "error",
        "progress": {"stage": "error", "message": "previous indexing failed"},
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", manager.base_dir)

    class _Signature:
        def hash(self) -> str:
            return "sig"

    embedding_signature = importlib.import_module("cognivoy.services.rag.embedding_signature")
    index_versioning = importlib.import_module("cognivoy.services.rag.index_versioning")
    monkeypatch.setattr(
        embedding_signature, "signature_from_embedding_config", lambda: _Signature()
    )
    monkeypatch.setattr(
        index_versioning,
        "find_matching_version",
        lambda *_args, **_kwargs: {"layout": "flat", "ready": True},
    )

    async def _noop_reindex_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_reindex_task", _noop_reindex_task)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/failed-kb/retry")

    assert response.status_code == 200
    body = response.json()
    assert body["noop"] is False
    assert isinstance(body.get("task_id"), str) and body["task_id"]
    assert manager.config["knowledge_bases"]["failed-kb"]["status"] == "initializing"


def test_retry_rejects_non_error_kb(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["ready-kb"] = {
        "path": "ready-kb",
        "status": "ready",
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", manager.base_dir)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/ready-kb/retry")

    assert response.status_code == 409
    assert "not in an error state" in response.json()["detail"]


def test_reindex_bypasses_existing_match_when_vectors_are_invalid(
    monkeypatch, tmp_path: Path
) -> None:
    monkeypatch.setenv("COGNIVOY_RAG_RETRIEVAL_PROFILE", "vector")
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    manager.config["knowledge_bases"]["bad-index-kb"] = {
        "path": "bad-index-kb",
        "status": "ready",
    }
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    monkeypatch.setattr(knowledge_router_module, "_kb_base_dir", manager.base_dir)

    class _Signature:
        def hash(self) -> str:
            return "sig"

    kb_dir = manager.base_dir / "bad-index-kb"
    version_dir = kb_dir / "version-1"
    version_dir.mkdir(parents=True)
    (version_dir / "docstore.json").write_text("{}", encoding="utf-8")
    (version_dir / "index_store.json").write_text("{}", encoding="utf-8")
    (version_dir / "meta.json").write_text(
        json.dumps({"signature": "sig", "version": "version-1"}),
        encoding="utf-8",
    )
    (version_dir / "default__vector_store.json").write_text(
        json.dumps({"embedding_dict": {"bad-node": [0.1, None, 0.3]}}),
        encoding="utf-8",
    )

    embedding_signature = importlib.import_module("cognivoy.services.rag.embedding_signature")
    monkeypatch.setattr(
        embedding_signature, "signature_from_embedding_config", lambda: _Signature()
    )

    async def _noop_reindex_task(*_args, **_kwargs):
        return None

    monkeypatch.setattr(knowledge_router_module, "run_reindex_task", _noop_reindex_task)

    with TestClient(_build_app()) as client:
        response = client.post("/api/v1/knowledge/bad-index-kb/reindex")

    assert response.status_code == 200
    body = response.json()
    assert body["noop"] is False
    assert isinstance(body.get("task_id"), str) and body["task_id"]
    assert manager.config["knowledge_bases"]["bad-index-kb"]["status"] == "initializing"


def test_update_config_coerces_legacy_provider_to_llamaindex() -> None:
    """Legacy `rag_provider` values are accepted and normalized to llamaindex."""

    class _FakeConfigService:
        def __init__(self) -> None:
            self.config: dict = {}

        def set_kb_config(self, kb_name: str, config: dict) -> None:
            self.kb_name = kb_name
            self.config = config

        def get_kb_config(self, _kb_name: str) -> dict:
            return self.config

    fake_service = _FakeConfigService()

    config_module = importlib.import_module("cognivoy.services.config")
    app = _build_app()

    with pytest.MonkeyPatch.context() as monkeypatch:
        monkeypatch.setattr(config_module, "get_kb_config_service", lambda: fake_service)
        with TestClient(app) as client:
            response = client.put(
                "/api/v1/knowledge/demo/config",
                json={"rag_provider": "raganything"},
            )

    assert response.status_code in {200, 204}
    assert fake_service.config.get("rag_provider") == "llamaindex"


def test_update_config_preserves_known_provider() -> None:
    class _FakeConfigService:
        def __init__(self) -> None:
            self.config: dict = {}

        def set_kb_config(self, kb_name: str, config: dict) -> None:
            self.kb_name = kb_name
            self.config = config

        def get_kb_config(self, _kb_name: str) -> dict:
            return self.config

    fake_service = _FakeConfigService()

    config_module = importlib.import_module("cognivoy.services.config")
    app = _build_app()

    with pytest.MonkeyPatch.context() as monkeypatch:
        monkeypatch.setattr(config_module, "get_kb_config_service", lambda: fake_service)
        with TestClient(app) as client:
            response = client.put(
                "/api/v1/knowledge/demo/config",
                json={"rag_provider": "pageindex"},
            )

    assert response.status_code in {200, 204}
    assert fake_service.config.get("rag_provider") == "pageindex"


def test_update_config_rejects_provider_change_for_ready_index(monkeypatch, tmp_path: Path) -> None:
    kb_dir = tmp_path / "demo"
    kb_dir.mkdir(parents=True)
    _write_ready_llamaindex_version(kb_dir)

    class _FakeConfigService:
        def __init__(self) -> None:
            self.config: dict = {"rag_provider": "llamaindex"}

        def set_kb_config(self, kb_name: str, config: dict) -> None:
            self.kb_name = kb_name
            self.config.update(config)

        def get_kb_config(self, _kb_name: str) -> dict:
            return dict(self.config)

    fake_service = _FakeConfigService()
    config_module = importlib.import_module("cognivoy.services.config")

    monkeypatch.setattr(config_module, "get_kb_config_service", lambda: fake_service)
    monkeypatch.setattr(knowledge_router_module, "_current_kb_base_dir", lambda: tmp_path)

    with TestClient(_build_app()) as client:
        response = client.put(
            "/api/v1/knowledge/demo/config",
            json={"rag_provider": "pageindex"},
        )

    assert response.status_code == 409
    assert "ready llamaindex index" in response.json()["detail"]
    assert fake_service.config["rag_provider"] == "llamaindex"


def test_rag_providers_marks_linkable() -> None:
    with TestClient(_build_app()) as client:
        providers = client.get("/api/v1/knowledge/rag-providers").json()["providers"]
    by_id = {p["id"]: p for p in providers}
    # Self-contained local indexes can be linked in place; PageIndex (cloud) and
    # LightRAG Server (remote, no local folder) can't.
    assert by_id["llamaindex"]["linkable"] is True
    assert by_id["graphrag"]["linkable"] is True
    assert by_id["lightrag"]["linkable"] is True
    assert by_id["pageindex"]["linkable"] is False
    assert by_id["lightrag-server"]["linkable"] is False


def test_probe_folder_endpoint_finds_ready_index(tmp_path: Path) -> None:
    version = tmp_path / "version-1"
    version.mkdir()
    (version / "docstore.json").write_text("{}", encoding="utf-8")
    (version / "index_store.json").write_text("{}", encoding="utf-8")
    (version / "meta.json").write_text(
        json.dumps({"version": "version-1", "signature": "x", "layout": "flat"}),
        encoding="utf-8",
    )

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/probe-folder",
            json={"folder_path": str(tmp_path), "rag_provider": "llamaindex"},
        )

    assert response.status_code == 200
    payload = response.json()
    assert payload["ok"] is True
    assert payload["version"] == "version-1"


def test_probe_folder_endpoint_rejects_pageindex(tmp_path: Path) -> None:
    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/probe-folder",
            json={"folder_path": str(tmp_path), "rag_provider": "pageindex"},
        )
    assert response.status_code == 200
    payload = response.json()
    assert payload["ok"] is False
    assert payload["error"]


def _patch_server_probe(monkeypatch, *, ok: bool, error: str | None = None) -> None:
    """Stub the LightRAG server probe so router tests need no live server."""
    from cognivoy.services.rag.pipelines.lightrag_server import probe as probe_module

    async def _fake_probe(server_url: str, api_key: str = "", **_kwargs):
        result = probe_module.ServerProbe(base_url=server_url.rstrip("/"))
        result.ok = ok
        result.reachable = ok
        result.auth_required = bool(api_key)
        result.auth_ok = ok
        result.error = error
        return result

    monkeypatch.setattr(probe_module, "probe_server", _fake_probe)


def test_probe_lightrag_server_endpoint_reports_verdict(monkeypatch) -> None:
    _patch_server_probe(monkeypatch, ok=True)
    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/probe-lightrag-server",
            json={"server_url": "http://localhost:9621", "api_key": "k"},
        )
    assert response.status_code == 200
    payload = response.json()
    assert payload["ok"] is True
    assert payload["base_url"] == "http://localhost:9621"


def test_connect_lightrag_server_registers_pointer(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    _patch_server_probe(monkeypatch, ok=True)

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/connect-lightrag-server",
            json={
                "name": "remote-kb",
                "server_url": "http://localhost:9621/",
                "api_key": "secret",
                "search_mode": "MIX",
            },
        )

    assert response.status_code == 200
    body = response.json()
    assert body["rag_provider"] == "lightrag-server"
    entry = manager.config["knowledge_bases"]["remote-kb"]
    assert entry["type"] == "lightrag_server"
    assert entry["server_url"] == "http://localhost:9621"
    assert entry["search_mode"] == "mix"  # normalized + validated


def test_connect_lightrag_server_rejects_unreachable(monkeypatch, tmp_path: Path) -> None:
    manager = _FakeKBManager(tmp_path / "knowledge_bases")
    monkeypatch.setattr(knowledge_router_module, "get_kb_manager", lambda: manager)
    _patch_server_probe(monkeypatch, ok=False, error="Could not reach a LightRAG server")

    with TestClient(_build_app()) as client:
        response = client.post(
            "/api/v1/knowledge/connect-lightrag-server",
            json={"name": "bad", "server_url": "http://nope:9621"},
        )

    assert response.status_code == 400
    assert "LightRAG" in response.json()["detail"]
    assert "bad" not in manager.config["knowledge_bases"]


def test_assert_not_connected_kb_blocks_connected_writes() -> None:
    from fastapi import HTTPException

    guard = knowledge_router_module._assert_not_connected_kb
    for kind in ("linked", "obsidian", "lightrag_server"):
        with pytest.raises(HTTPException) as excinfo:
            guard("kb", {"type": kind})
        assert excinfo.value.status_code == 409
    # An ordinary KB is writable — the guard is a no-op.
    guard("kb", {"path": "kb", "status": "ready"})
